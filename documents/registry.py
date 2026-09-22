"""Реєстр ТТН за період - замінює макет реальною
агрегацією по базі."""
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import TTN


def build_monthly_registry_workbook(start_date, end_date, counterparty=None, vehicle=None):
    qs = TTN.objects.filter(date__range=(start_date, end_date)).select_related(
        'trip', 'trip__contract__counterparty', 'trip__vehicle', 'trip__driver',
    ).order_by('date')
    if counterparty is not None:
        qs = qs.filter(trip__contract__counterparty=counterparty)
    if vehicle is not None:
        qs = qs.filter(trip__vehicle=vehicle)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Реєстр ТТН'
    headers = ['Дата', 'Номер ТТН', 'Контрагент', 'Автомобіль', 'Водій', 'Вага, т', 'Сума, грн']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    row_idx = 1
    for ttn in qs:
        trip = ttn.trip
        row_idx += 1
        ws.append([
            trip.date, ttn.number, trip.contract.counterparty.name,
            trip.vehicle.plate_number, trip.driver.full_name,
            float(trip.weight_t), float(trip.amount_no_vat),
        ])

    if row_idx > 1:
        total_row = row_idx + 1
        ws.cell(row=total_row, column=1, value='Разом:').font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=f'=SUM(F2:F{row_idx})').font = Font(bold=True)
        ws.cell(row=total_row, column=7, value=f'=SUM(G2:G{row_idx})').font = Font(bold=True)

    widths = [12, 14, 32, 14, 24, 10, 14]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return wb
