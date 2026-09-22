"""Зведений звіт по контрагенту/автопарку/водію за період -
кількість рейсів і суми без ПДВ/ПДВ/з ПДВ, основа для акту звірки."""
from decimal import Decimal

import openpyxl
from django.db.models import Count, Sum
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.vat import calculate_vat
from trips.models import Trip

GROUP_FIELDS = {
    'counterparty': ('contract__counterparty_id', 'contract__counterparty__name', 'Контрагент'),
    'vehicle': ('vehicle_id', 'vehicle__plate_number', 'Автомобіль'),
    'driver': ('driver_id', 'driver__full_name', 'Водій'),
}


def build_summary(start_date, end_date, group_by='counterparty'):
    id_field, name_field, _ = GROUP_FIELDS[group_by]
    qs = Trip.objects.filter(
        date__range=(start_date, end_date), status=Trip.Status.CONFIRMED,
    ).values(id_field, name_field).annotate(
        trips=Count('id'), amount_no_vat=Sum('amount_no_vat'),
    ).order_by('-amount_no_vat')

    result = []
    for row in qs:
        amount_no_vat = row['amount_no_vat'] or Decimal('0')
        vat_amount, amount_with_vat = calculate_vat(amount_no_vat)
        result.append({
            'name': row[name_field],
            'trips': row['trips'],
            'amount_no_vat': amount_no_vat,
            'vat_amount': vat_amount,
            'amount_with_vat': amount_with_vat,
        })
    return result


def build_summary_workbook(start_date, end_date, group_by='counterparty'):
    _, _, label = GROUP_FIELDS[group_by]
    rows = build_summary(start_date, end_date, group_by)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Зведений звіт'
    headers = [label, 'Кількість рейсів', 'Сума без ПДВ, грн', 'ПДВ, грн', 'Сума з ПДВ, грн']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([
            row['name'], row['trips'], float(row['amount_no_vat']),
            float(row['vat_amount']), float(row['amount_with_vat']),
        ])

    last_row = len(rows) + 1
    if rows:
        total_row = last_row + 1
        ws.cell(row=total_row, column=1, value='Разом:').font = Font(bold=True)
        ws.cell(row=total_row, column=2, value=f'=SUM(B2:B{last_row})').font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=f'=SUM(C2:C{last_row})').font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=f'=SUM(D2:D{last_row})').font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=f'=SUM(E2:E{last_row})').font = Font(bold=True)

    widths = [32, 16, 18, 14, 18]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return wb
