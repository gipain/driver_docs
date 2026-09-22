"""fill_ttn(ttn) - заповнює заготовку ТТН.xlsx даними конкретного
рейсу за зафіксованими координатами з documents/ttn_cells.py.
Повертає openpyxl Workbook - виклик мoже або зберегти на диск, або
віддати як .xlsx-відповідь, або передати далі на конвертацію в PDF."""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

from core.number_to_words import count_to_words_uk, fractional_to_words_uk, money_to_words_uk
from core.vat import calculate_vat
from documents.ttn_cells import (
    TTN_CARGO_LINE_COLUMNS,
    TTN_CARGO_LINE_ROW,
    TTN_CELLS,
    TTN_SIGNATORY_CELLS,
    TTN_TOTAL_ROW,
)

TEMPLATE_PATH = Path(__file__).resolve().parent / 'xlsx_templates' / 'ttn_template.xlsx'

UA_MONTHS_GENITIVE = [
    'Січня', 'Лютого', 'Березня', 'Квітня', 'Травня', 'Червня',
    'Липня', 'Серпня', 'Вересня', 'Жовтня', 'Листопада', 'Грудня',
]


def _date_words(d) -> str:
    """«4 Вересня 2026 р.» - точно у форматі наданого зразка ТТН."""
    return f'{d.day} {UA_MONTHS_GENITIVE[d.month - 1]} {d.year} р.'


def fill_ttn(ttn) -> openpyxl.Workbook:
    trip = ttn.trip
    vehicle = trip.vehicle
    vat_amount, amount_with_vat = calculate_vat(trip.amount_no_vat)

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    values = {
        'header_line': f'№ {ttn.number} {_date_words(ttn.date)}',
        'place_of_drafting': ttn.place,
        'vehicle_line': f'{vehicle.brand_model}, {vehicle.vehicle_type}, держ.№ {vehicle.plate_number}',
        'trailer_line': trip.trailer.plate_number if trip.trailer else '',
        'transport_type': vehicle.vehicle_type,
        'carrier_address': trip.carrier.address,
        'carrier_line': f'{trip.carrier.name}, {trip.carrier.edrpou}',
        'driver_line': f'{trip.driver.full_name}, {trip.driver.license_number}',
        'sender_line': f'{trip.sender.name}, {trip.sender.edrpou}',
        'receiver_line': f'{trip.receiver.name}, {trip.receiver.edrpou}',
        'loading_address': trip.loading_address,
        'unloading_address': trip.unloading_address,
        'places_words': count_to_words_uk(trip.places_count, 'місце', 'місця', 'місць'),
        'weight_words': fractional_to_words_uk(trip.weight_t),
        'driver_signature_name': trip.driver.full_name,
        'vehicle_length': float(vehicle.length_m),
        'vehicle_width': float(vehicle.width_m),
        'vehicle_height': float(vehicle.height_m),
        'vehicle_total_weight': float(vehicle.capacity_t),
        'amount_words': money_to_words_uk(amount_with_vat),
        'vat_line': f'{vat_amount} грн.',
        'accompanying_docs': ttn.accompanying_docs,
    }
    for field, coord in TTN_CELLS.items():
        ws[coord] = values[field]

    row = TTN_CARGO_LINE_ROW
    for line in ttn.cargo_lines.all():
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_name']}{row}"] = line.name
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_unit']}{row}"] = line.unit
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_qty']}{row}"] = line.qty
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_unit_price']}{row}"] = float(line.unit_price)
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_amount']}{row}"] = float(line.amount)
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_packaging']}{row}"] = line.packaging
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_docs']}{row}"] = ttn.accompanying_docs
        ws[f"{TTN_CARGO_LINE_COLUMNS['cargo_weight']}{row}"] = float(line.weight_t)
        row += 1

    # Підсумковий рядок "Всього" - формули SUM (день 10), а не готові числа,
    # щоб перераховувались автоматично, якщо рядки зміняться вручну.
    first_row, last_row = TTN_CARGO_LINE_ROW, row - 1
    ws[f"Z{TTN_TOTAL_ROW}"] = f'=SUM(Z{first_row}:Z{last_row})'
    ws[f"AH{TTN_TOTAL_ROW}"] = f'=SUM(AH{first_row}:AH{last_row})'
    ws[f"AW{TTN_TOTAL_ROW}"] = f'=SUM(AW{first_row}:AW{last_row})'

    ws[TTN_SIGNATORY_CELLS['sender_signatory']] = f'{trip.sender.director}, Директор'
    ws[TTN_SIGNATORY_CELLS['loading_weight']] = float(trip.weight_t)

    _configure_print_layout(ws)
    return wb


def _configure_print_layout(ws) -> None:
    """Наданий зразок і сама заготовка без цих налаштувань друкуються
    LibreOffice на 6-8 фрагментованих сторінок (лист дуже широкий - 52
    колонки). Друкований лист має лишатись бухгалтерськи звичним - лицьовий
    бік (рядки 1-24) і зворотний (рядки 26-50) окремими сторінками A4."""
    from openpyxl.worksheet.pagebreak import ColBreak, RowBreak

    # A20 успадкувала з наданого зразка перенесення тексту (wrapText) у
    # колонці шириною ~2 символи - підпис "Відомості про транспортний
    # засіб..." розбивається по одній літері на рядок і роздуває висоту
    # рядка до 409 пунктів, через що текст виглядає нечитабельним
    # вертикальним стовпчиком, а дані габаритів авто зʼїжджають у самий
    # низ сторінки. Саме вимкнення wrapText цього не виправляє: усередині
    # значення є жорсткий перенос рядка (\n), який Excel/LibreOffice
    # показує окремим рядком незалежно від wrapText, а рядок 20 замалий
    # для двох рядків - другий рядок тексту просто обрізається по висоті.
    # Виправляємо все разом: власна об'єднана комірка A20:R20 (ці стовпці
    # порожні в цьому рядку - значення габаритів пишуться від S20 і
    # правіше, і межа R обрана саме перед ними), жорсткий перенос
    # замінено на пробіл, а шрифт зменшено до 6пт (як у сусідніх підписах
    # на кшталт AO21) - інакше навіть у цій ширшій комірці рядок не
    # влазить і обрізається за кілька символів до кінця.
    a20 = ws['A20']
    a20.value = a20.value.replace('\n', ' ')
    ws.merge_cells('A20:R20')
    a20.alignment = Alignment(wrap_text=False, vertical='center')
    a20_font = copy(a20.font)
    a20_font.size = 6
    a20.font = a20_font
    ws.row_dimensions[20].height = 20

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.scale = 40
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_area = 'A1:AZ24,A26:AZ50'
    # Заготовка успадкувала ручні розриви сторінки з наданого зразка -
    # вони конфліктують із власним масштабуванням і ріжуть аркуш на зайві
    # фрагменти, тому прибираємо їх перед друком.
    ws.row_breaks = RowBreak()
    ws.col_breaks = ColBreak()
